"""Turn uploaded files into plain text for chunking and embedding.

Supported: .txt .md (read as-is), .pdf (pypdf), .xlsx/.xlsm (openpyxl), .csv.

Tabular data is rendered one row per line with the column name repeated on every
value ("Region: North | Sales: 1200"). That is deliberate: a chunk of raw CSV
rows loses its header once it is split, so a retrieved chunk would be numbers
with no meaning. Repeating the column names keeps every row self-describing.
"""

import csv
import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


class ExtractionError(Exception):
    """The file could not be turned into usable text."""


# Extension -> human label, also the single source of truth for what we accept.
SUPPORTED = {
    ".txt": "text",
    ".md": "markdown",
    ".pdf": "PDF",
    ".xlsx": "Excel",
    ".xlsm": "Excel",
    ".csv": "CSV",
}

MAX_CELL = 300        # truncate absurdly long cells
MAX_ROWS = 20_000     # guard against a runaway spreadsheet


def _clean(value) -> str:
    """Render one cell as a short, single-line string."""
    if value is None:
        return ""
    text = str(value).strip().replace("\n", " ").replace("\r", " ")
    return text[:MAX_CELL] + "…" if len(text) > MAX_CELL else text


def _rows_to_text(rows, sheet_name: str | None = None) -> str:
    """Render rows (first row = header) as one self-describing line per row."""
    rows = [r for r in rows if any(_clean(c) for c in r)]  # drop blank rows
    if not rows:
        return ""

    header = [_clean(c) or f"col{i + 1}" for i, c in enumerate(rows[0])]
    lines = []
    if sheet_name:
        lines.append(f"Sheet: {sheet_name}")
    lines.append("Columns: " + ", ".join(header))

    for row in rows[1:]:
        pairs = [
            f"{header[i] if i < len(header) else f'col{i + 1}'}: {_clean(cell)}"
            for i, cell in enumerate(row)
            if _clean(cell)
        ]
        if pairs:
            prefix = f"[{sheet_name}] " if sheet_name else ""
            lines.append(prefix + " | ".join(pairs))
    return "\n".join(lines)


def _extract_pdf(data: bytes) -> str:
    """Pull the text layer out of a PDF, page by page."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError("pypdf is not installed. Run: pip install -r requirements.txt") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not open the PDF: {exc}") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")           # many PDFs are encrypted with an empty password
        except Exception as exc:
            raise ExtractionError("The PDF is password-protected.") from exc

    pages = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            logger.warning("Could not extract page %d", number, exc_info=True)
            continue
        if text:
            # Page markers help you trace an answer back to a location.
            pages.append(f"[page {number}]\n{text}")

    if not pages:
        raise ExtractionError(
            "No text found in this PDF. Scanned or image-only PDFs need OCR, which is not built in."
        )
    return "\n\n".join(pages)


def _extract_excel(data: bytes) -> str:
    """Read every sheet of an .xlsx/.xlsm workbook."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ExtractionError("openpyxl is not installed. Run: pip install -r requirements.txt") from exc

    try:
        # read_only streams rows; data_only gives computed values instead of formulas.
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError(
            f"Could not open the workbook: {exc}. Legacy .xls is not supported — save it as .xlsx."
        ) from exc

    try:
        blocks = []
        for sheet in workbook.worksheets:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS:
                    logger.warning("Sheet %r truncated at %d rows", sheet.title, MAX_ROWS)
                    break
                rows.append(row)
            block = _rows_to_text(rows, sheet_name=sheet.title)
            if block:
                blocks.append(block)
    finally:
        workbook.close()

    if not blocks:
        raise ExtractionError("The workbook has no readable data.")
    return "\n\n".join(blocks)


def _extract_csv(data: bytes) -> str:
    """Read a CSV, sniffing the delimiter (comma, semicolon, tab…)."""
    text = data.decode("utf-8", errors="replace")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel                       # fall back to plain comma

    try:
        rows = list(csv.reader(io.StringIO(text), dialect))[:MAX_ROWS]
    except csv.Error as exc:
        raise ExtractionError(f"Could not parse the CSV: {exc}") from exc

    block = _rows_to_text(rows)
    if not block:
        raise ExtractionError("The CSV has no readable rows.")
    return block


def _extract_plain(data: bytes) -> str:
    """Decode a .txt/.md file, tolerating the odd bad byte."""
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        logger.info("File is not valid UTF-8; decoding with replacement")
        return data.decode("utf-8", errors="replace")


def extract_text(data: bytes, filename: str) -> str:
    """Extract plain text from `data`, choosing the reader by file extension.

    Raises ExtractionError with a message suitable for showing to the user.
    """
    if not data:
        raise ExtractionError("The file is empty.")

    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        supported = ", ".join(sorted(SUPPORTED))
        raise ExtractionError(f"{suffix or 'this file type'} is not supported. Supported: {supported}")

    try:
        if suffix == ".pdf":
            text = _extract_pdf(data)
        elif suffix in {".xlsx", ".xlsm"}:
            text = _extract_excel(data)
        elif suffix == ".csv":
            text = _extract_csv(data)
        else:
            text = _extract_plain(data)
    except ExtractionError:
        raise                                     # already has a good message
    except Exception as exc:
        logger.exception("Unexpected failure extracting %s", filename)
        raise ExtractionError(f"Could not read {filename}: {exc}") from exc

    if not text.strip():
        raise ExtractionError("No readable text was found in this file.")
    return text
